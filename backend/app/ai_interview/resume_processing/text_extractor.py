import io
import fitz  # PyMuPDF
from docx import Document
from app.ai_interview.resume_processing.schemas import ResumeFileInput, ExtractedDocument
from app.ai_interview.resume_processing.exceptions import (
    CorruptedFileError,
    UnsupportedFileTypeError
)
import os

class TextExtractor:
    @staticmethod
    def extract(file_input: ResumeFileInput) -> ExtractedDocument:
        ext = os.path.splitext(file_input.filename)[1].lower()
        
        raw_text = ""
        page_count = 0
        
        if ext == ".pdf":
            raw_text, page_count = TextExtractor._extract_from_pdf(file_input.file_bytes)
        elif ext in [".doc", ".docx"]:
            raw_text, page_count = TextExtractor._extract_from_docx(file_input.file_bytes)
        else:
            raise UnsupportedFileTypeError(f"Cannot extract from unsupported extension: {ext}")
            
        return ExtractedDocument(
            raw_text=raw_text,
            page_count=page_count,
            character_count=len(raw_text),
            source_type=ext[1:]  # e.g., 'pdf', 'docx'
        )
        
    @staticmethod
    def _extract_from_pdf(file_bytes: bytes) -> tuple[str, int]:
        text = ""
        page_count = 0
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            page_count = len(doc)
            for page in doc:
                page_text = page.get_text()
                if page_text:
                    text += page_text + "\n"
        except Exception as e:
            raise CorruptedFileError(f"Failed to open/parse PDF document: {e}")
            
        return text, page_count

    @staticmethod
    def _extract_from_docx(file_bytes: bytes) -> tuple[str, int]:
        text = ""
        try:
            doc = Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs:
                if para.text:
                    text += para.text + "\n"
            page_count = 1 # DOCX doesn't have a reliable page count without rendering
        except Exception as e:
            raise CorruptedFileError(f"Failed to open/parse DOCX document: {e}")
            
        return text, page_count
